"""PowerPoint template -> Beamer theme, and presentations in the document tab."""
import re
import zipfile

import pytest

from NoorSuite import latex, pptximport
from NoorSuite import templates as T
from tests.test_gui_features import app, win, _doc_win  # noqa: F401  (window fixtures)

EMU = 360000


def _png_bytes(w=40, h=20, color=0xFFB02030):
    from PyQt6.QtCore import QBuffer, QByteArray, QIODevice
    from PyQt6.QtGui import QGuiApplication, QImage
    QGuiApplication.instance() or QGuiApplication([])
    img = QImage(w, h, QImage.Format.Format_ARGB32)
    img.fill(color)
    ba = QByteArray()
    buf = QBuffer(ba)
    buf.open(QIODevice.OpenModeFlag.WriteOnly)
    img.save(buf, "PNG")
    return bytes(ba)


def _sp(name, x, y, w, h, fill=None, ph=None, text=None, size=None, color=None):
    ph_xml = f'<p:nvPr><p:ph type="{ph[0]}"{" idx=%r" % ph[1] if ph[1] else ""}/></p:nvPr>' \
        if ph else "<p:nvPr/>"
    ph_xml = ph_xml.replace("'", '"')
    fill_xml = f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill>' if fill else ""
    rpr = (f'<a:rPr lang="en-US" sz="{int(size * 100)}">'
           + (f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>' if color else "")
           + "</a:rPr>") if size else ""
    body = f"<a:p><a:r>{rpr}<a:t>{text}</a:t></a:r></a:p>" if text else ""
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="2" name="{name}"/><p:cNvSpPr/>{ph_xml}</p:nvSpPr>'
            f'<p:spPr><a:xfrm><a:off x="{int(x * EMU)}" y="{int(y * EMU)}"/>'
            f'<a:ext cx="{int(w * EMU)}" cy="{int(h * EMU)}"/></a:xfrm>{fill_xml}</p:spPr>'
            f"<p:txBody><a:bodyPr/>{body}</p:txBody></p:sp>")


def _pic(name, rid, x, y, w, h):
    return (f'<p:pic><p:nvPicPr><p:cNvPr id="3" name="{name}"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>'
            f'<p:blipFill><a:blip r:embed="{rid}"/></p:blipFill><p:spPr><a:xfrm>'
            f'<a:off x="{int(x * EMU)}" y="{int(y * EMU)}"/><a:ext cx="{int(w * EMU)}" '
            f'cy="{int(h * EMU)}"/></a:xfrm></p:spPr></p:pic>')


def _make_pptx(path):
    ns = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
          'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
    theme = ('<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:themeElements>'
             '<a:clrScheme name="x"><a:dk1><a:srgbClr val="000000"/></a:dk1><a:lt1><a:srgbClr '
             'val="FFFFFF"/></a:lt1><a:accent1><a:srgbClr val="336699"/></a:accent1></a:clrScheme>'
             '<a:fontScheme><a:majorFont><a:latin typeface="Open Sans Light"/></a:majorFont>'
             '<a:minorFont><a:latin typeface="Open Sans Light"/></a:minorFont></a:fontScheme>'
             "</a:themeElements></a:theme>")
    master = (f"<p:sldMaster {ns}><p:cSld><p:spTree>"
              + _sp("Bar", 0, 0, 33.867, 2.5, fill="BE1E3C")
              + _sp("Title", 0.7, 0.4, 30.7, 1.6, ph=("title", ""))
              + _sp("Body", 0.7, 3.2, 32.5, 15.3, ph=("body", "1"))
              + _pic("Logo", "rId2", 0.2, 18.9, 5.5, 2.1)
              + _sp("Num", 31.8, 0.6, 1.4, 1.0, ph=("sldNum", ""), text="1", size=20, color="FFFFFF")
              + "</p:spTree></p:cSld><p:txStyles><p:titleStyle><a:lvl1pPr><a:defRPr sz=\"3600\">"
              "<a:solidFill><a:srgbClr val=\"FFFFFF\"/></a:solidFill></a:defRPr></a:lvl1pPr>"
              "</p:titleStyle><p:bodyStyle><a:lvl1pPr><a:defRPr sz=\"2400\"/></a:lvl1pPr>"
              "</p:bodyStyle></p:txStyles></p:sldMaster>")
    layout = (f'<p:sldLayout {ns} showMasterSp="0"><p:cSld name="Title"><p:spTree>'
              + _pic("Photo", "rId2", 0, 0, 33.867, 14.4)
              + _sp("Title", 0.7, 14.7, 26.9, 3.2, ph=("body", "13"), text="x", size=38,
                    color="0069B3")
              + _sp("Name", 0.7, 17.9, 26.9, 1.0, ph=("body", "14"), text="x", size=20)
              + _sp("Date", 0.7, 18.9, 26.9, 1.8, ph=("body", "15"), text="x", size=18)
              + _sp("Credit", 27.8, 13.5, 6.1, 0.9, fill="FFFFFF", text="Photo: Someone", size=14)
              + "</p:spTree></p:cSld></p:sldLayout>")
    content = f"<p:sldLayout {ns}><p:cSld name=\"Content\"><p:spTree/></p:cSld></p:sldLayout>"
    slide1 = f"<p:sld {ns}><p:cSld><p:spTree/></p:cSld></p:sld>"
    rel = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
           '{}</Relationships>')
    png = _png_bytes()
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("ppt/presentation.xml",
                   f'<p:presentation {ns}><p:sldSz cx="{int(33.867 * EMU)}" cy="{int(21.167 * EMU)}"/>'
                   "</p:presentation>")
        z.writestr("ppt/theme/theme1.xml", theme)
        z.writestr("ppt/slideMasters/slideMaster1.xml", master)
        z.writestr("ppt/slideMasters/_rels/slideMaster1.xml.rels", rel.format(
            '<Relationship Id="rId2" Target="../media/logo.png"/>'))
        z.writestr("ppt/slideLayouts/slideLayout1.xml", layout)
        z.writestr("ppt/slideLayouts/_rels/slideLayout1.xml.rels", rel.format(
            '<Relationship Id="rId2" Target="../media/photo.png"/>'))
        z.writestr("ppt/slideLayouts/slideLayout2.xml", content)
        z.writestr("ppt/slides/slide1.xml", slide1)
        z.writestr("ppt/slides/_rels/slide1.xml.rels", rel.format(
            '<Relationship Id="rId1" Target="../slideLayouts/slideLayout1.xml"/>'))
        z.writestr("ppt/media/logo.png", png)
        z.writestr("ppt/media/photo.png", _png_bytes(80, 30, 0xFF2060A0))


@pytest.fixture
def pres_env(tmp_path, monkeypatch):
    monkeypatch.setenv("NOORSUITE_PRESENTATIONS", str(tmp_path / "pres"))
    monkeypatch.setenv("NOORSUITE_TEMPLATES", str(tmp_path / "tpl"))
    p = tmp_path / "Corporate Template.pptx"
    _make_pptx(p)
    return p


def test_pptx_geometry_is_read(pres_env):
    P = pptximport.parse(pres_env)
    assert abs(P["size_cm"][0] - 33.867) < 0.01
    bars = [i for i in P["master"]["items"] if i.get("fill") == "BE1E3C"]
    assert bars and abs(bars[0]["h"] - 2.5) < 1e-6
    assert P["title_style"] == {"size": 36.0, "color": "FFFFFF"}
    assert P["fonts"][0] == "Open Sans Light"
    lay = P["layouts"]["ppt/slideLayouts/slideLayout1.xml"]
    assert lay["hides_master"] and lay["name"] == "Title"


def test_import_writes_theme_media_and_a_beamer_layout(pres_env, tmp_path):
    res = pptximport.import_pptx(pres_env)
    assert res["theme"] == "corporatetemplate" and res["id"] == "user:corporatetemplate"
    d = tmp_path / "pres" / "corporatetemplate"
    sty = (d / "beamerthemecorporatetemplate.sty").read_text(encoding="utf8")
    assert r"\fill[corporatetemplatec0]" in sty and "BE1E3C" in sty        # the red bar
    assert "Photo: Someone" in sty and r"\inserttitle" in sty and r"\insertdate" in sty
    assert (d / "corporatetemplate-logo.jpg").is_file() or \
        (d / "corporatetemplate-logo.png").is_file()
    # the layout is listed as a presentation layout of the beamer base
    assert res["id"] in dict(T.list_templates()) and T.base_of(res["id"]) == "beamer"
    doc = T.build_document(res["id"], "My talk", ["A. One"], "", "")
    assert r"\documentclass[aspectratio=1610,11pt,t]{beamer}" in doc
    assert r"\usetheme{corporatetemplate}" in doc and r"\titlepage" in doc
    assert T.detect(doc) == "beamer"


def test_theme_files_travel_with_the_document(pres_env, tmp_path):
    res = pptximport.import_pptx(pres_env)
    doc = T.build_document(res["id"], "T", ["A"], "", "")
    files = pptximport.support_files(doc)
    assert "beamerthemecorporatetemplate.sty" in files and len(files) >= 3
    wd = latex.prepare_workdir(tmp_path / "wd", doc)
    assert (wd / "beamerthemecorporatetemplate.sty").is_file()
    assert pptximport.support_files(r"\usetheme{Madrid}") == {}


@pytest.mark.skipif(latex.find_engine() is None, reason="no LaTeX engine installed")
def test_generated_theme_compiles(pres_env, tmp_path):
    res = pptximport.import_pptx(pres_env)
    body = (r"\begin{frame}{A frame}" "\n" r"\begin{itemize}\item One\item Two\end{itemize}" "\n"
            r"\framesource{Source: test}" "\n" r"\end{frame}" "\n")
    doc = T.build_document(res["id"], "My talk", ["A. One"], "", body)
    latex.prepare_workdir(tmp_path / "w", doc)
    out = latex.build_pdf(tmp_path / "w")
    assert out.ok, out.log[-800:]


def test_beamer_switching_keeps_frames_and_does_not_double_the_title(pres_env):
    res = pptximport.import_pptx(pres_env)
    plain = T.build_document("beamer", "T", ["A"], "", r"\begin{frame}{X}Hi\end{frame}")
    mine = T.apply_template(plain, res["id"])
    assert mine.count(r"\titlepage") == 1 and r"\begin{frame}{X}" in mine
    assert r"\usetheme{corporatetemplate}" in mine
    back = T.apply_template(mine, "beamer")
    assert back.count(r"\titlepage") == 1 and r"\usetheme{default}" in back
    # citations in a presentation do not get natbib injected
    assert not latex.own_citations(plain)


def test_figure_frames_and_layout_choices_in_the_document_tab(win, pres_env):
    res = pptximport.import_pptx(pres_env)
    ids, dm, tab = _doc_win(win, 1)
    tab.set_source(T.build_document(res["id"], "Talk", ["A"], "", r"\begin{frame}{X}Hi\end{frame}"))
    tab.model.layout = res["id"]
    assert tab.current_layout() == res["id"]
    labels = [tab.layout_combo.itemData(i) for i in range(tab.layout_combo.count())]
    assert "beamer" in labels and res["id"] in labels and "acs" not in labels
    from PyQt6.QtCore import Qt
    fp = tab.figures
    for k in range(fp.list.count()):
        if fp.list.item(k).data(Qt.ItemDataRole.UserRole) == ids[0]:
            fp.list.item(k).setSelected(True)
    snip = fp.snippets_for(fp.list.selectedItems())
    assert snip.startswith(r"\begin{frame}{") and r"\noorsheet[width=0.85\textwidth]{" in snip
    assert r"\begin{figure}" not in snip
    # an article gets article layouts, not presentation ones
    tab.set_source(T.build_document("article", "T", ["A"], "A.", "B."))
    tab.sync_layout_combo()
    ids2 = [tab.layout_combo.itemData(i) for i in range(tab.layout_combo.count())]
    assert "acs" in ids2 and "beamer" not in ids2


def test_new_presentation_from_the_window(win, pres_env):
    res = pptximport.import_pptx(pres_env)
    dm = win.create_new_document("Group talk", template=res["id"])
    assert r"\usetheme{corporatetemplate}" in dm.source and "Group talk" in dm.source
    plain = win.create_new_document("Plain talk", template="beamer")
    assert r"\documentclass[aspectratio=1610,11pt,t]{beamer}" in plain.source


# ------------------------------------------------------------------ bundled neutral theme + bundles
def test_neutral_theme_ships_with_noorsuite_and_has_no_corporate_artwork():
    from pathlib import Path
    d = pptximport.bundled_themes_dir() / "noorclean"
    assert (d / "beamerthemenoorclean.sty").is_file()
    assert [f.name for f in d.iterdir()] == ["beamerthemenoorclean.sty"]     # text only: no logos
    assert "beamer-clean" in dict(T.list_templates()) and T.base_of("beamer-clean") == "beamer"
    doc = T.build_document("beamer-clean", "Talk", ["A"], "", "")
    assert r"\usetheme{noorclean}" in doc and T.detect(doc) == "beamer-clean"
    assert list(pptximport.support_files(doc)) == ["beamerthemenoorclean.sty"]


@pytest.mark.skipif(latex.find_engine() is None, reason="no LaTeX engine installed")
def test_neutral_theme_compiles_with_figures_and_a_source_line(tmp_path):
    body = (r"\begin{frame}{Results}" "\n" r"\begin{itemize}\item One\item Two\end{itemize}" "\n"
            r"\framesource{Source: measurement series 3}" "\n" r"\end{frame}" "\n")
    doc = T.build_document("beamer-clean", "A talk", ["A. One"], "", body)
    latex.prepare_workdir(tmp_path / "w", doc)
    out = latex.build_pdf(tmp_path / "w")
    assert out.ok, out.log[-800:]


def test_theme_bundle_roundtrip_moves_a_template_between_machines(pres_env, tmp_path, monkeypatch):
    res = pptximport.import_pptx(pres_env)
    bundle = pptximport.export_bundle(res["theme"], tmp_path / "talk.noortheme")
    assert bundle.is_file() and zipfile.is_zipfile(bundle)
    names = zipfile.ZipFile(bundle).namelist()
    assert "bundle.json" in names and "layout.json" in names
    assert any(n.endswith("beamertheme" + res["theme"] + ".sty") for n in names)
    # a second "machine": empty folders
    monkeypatch.setenv("NOORSUITE_PRESENTATIONS", str(tmp_path / "other_pres"))
    monkeypatch.setenv("NOORSUITE_TEMPLATES", str(tmp_path / "other_tpl"))
    assert res["id"] not in dict(T.list_templates())
    back = pptximport.import_bundle(bundle)
    assert back["id"] == res["id"] and res["id"] in dict(T.list_templates())
    doc = T.build_document(back["id"], "T", ["A"], "", "")
    assert "beamerthemecorporatetemplate.sty" in pptximport.support_files(doc)
    with pytest.raises(ValueError):
        pptximport.export_bundle("does-not-exist", tmp_path / "x.noortheme")


# ------------------------------------------------------------------ PowerPoint export + slide show
_TALK = (r"\begin{frame}{Results}{A subtitle}" "\n"
         r"  \begin{itemize}" "\n"
         r"    \item Level one with $E = mc^2$ and \textbf{bold}" "\n"
         r"      \begin{itemize}\item Second level $\sigma$ in S\,cm$^{-1}$\end{itemize}" "\n"
         r"    \item Another point" "\n"
         r"  \end{itemize}" "\n"
         r"  \framesource{Source: own data}" "\n" r"  \note{Say something about T.}" "\n"
         r"\end{frame}" "\n"
         r"\begin{frame}{A figure}\centering\noorsheet[width=0.8\textwidth]{fig-a}\end{frame}" "\n"
         r"\begin{frame}{Two columns}\begin{columns}" "\n"
         r"\begin{column}{0.5\textwidth}\begin{itemize}\item L1\item L2\end{itemize}\end{column}" "\n"
         r"\begin{column}{0.5\textwidth}\noorsheet{fig-a}\end{column}\end{columns}\end{frame}" "\n"
         r"\begin{frame}{Table}\begin{tabular}{rr}\toprule $T$ & $\sigma$ \\ \midrule 1 & 2 \\"
         r" \bottomrule\end{tabular}\end{frame}" "\n")


def test_math_and_markup_become_readable_slide_text():
    from NoorSuite import pptxexport as X
    assert X.plain(r"$\alpha^2 + \beta_i$") == "α² + βᵢ"
    assert X.plain(r"S\,cm$^{-1}$ and $\frac{a}{b}$") == "S cm⁻¹ and (a)/(b)"
    assert X.plain(r"\textbf{bold} \emph{it} 50\% --- done") == "bold it 50% — done"
    assert X.plain(r"Li$_{3}$PS$_{4}$") == "Li₃PS₄"


def test_beamer_frames_are_parsed_into_blocks():
    from NoorSuite import pptxexport as X
    src = T.build_document("beamer", "My talk", ["A. One"], "", _TALK).replace(
        r"\date{\today}", r"\date{1 Jan}")
    P = X.parse_presentation(src)
    assert P["meta"]["title"] == "My talk" and P["meta"]["author"] == "A. One"
    kinds = [f["kind"] for f in P["frames"]]
    assert kinds == ["title", "content", "content", "content", "content"]
    f1 = P["frames"][1]
    assert f1["title"] == "Results" and f1["subtitle"] == "A subtitle"
    lst = f1["blocks"][0][1]
    assert [(lvl, txt) for lvl, txt, _o in lst] == [
        (0, "Level one with E = mc² and bold"), (1, "Second level σ in S cm⁻¹"),
        (0, "Another point")]
    assert f1["source"] == "Source: own data" and f1["note"] == "Say something about T."
    assert P["frames"][2]["blocks"][0] == ("picture", "noorsheet", "fig-a", "[width=0.8\\textwidth]")
    assert P["frames"][3]["blocks"][0][0] == "columns" and len(P["frames"][3]["blocks"][0][1]) == 2
    assert P["frames"][4]["blocks"][0][0] == "table"


def _sheet_stub(tmp_path):
    from NoorSuite.model import SheetModel
    sm = SheetModel("Fig A", sheet_id="fig-a")
    png = _png_bytes(60, 40, 0xFF336699)

    def sheet_png(_sm, path):
        open(path, "wb").write(png)
    return [sm], sheet_png


def test_pptx_export_default_design_is_editable(tmp_path):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from NoorSuite import pptxexport as X
    sheets, sheet_png = _sheet_stub(tmp_path)
    src = T.build_document("beamer", "My talk", ["A. One"], "", _TALK)
    out = tmp_path / "talk.pptx"
    ok, notes = X.export_pptx(src, out, sheets, [], sheet_png, lambda *a: None, tmp_path / "w")
    assert ok and out.is_file()
    prs = Presentation(str(out))
    assert prs.slide_width / prs.slide_height == pytest.approx(1.6, abs=0.01)
    slides = list(prs.slides)
    assert len(slides) == 5
    title = [s.text_frame.text for s in slides[0].shapes if s.has_text_frame]
    assert "My talk" in title[0] and any("A. One" in t for t in title)
    body = next(s for s in slides[1].shapes if s.has_text_frame and "Level one" in s.text_frame.text)
    paras = body.text_frame.paragraphs
    assert [p.level for p in paras] == [0, 1, 0]
    assert slides[1].notes_slide.notes_text_frame.text == "Say something about T."
    assert any(s.shape_type == 13 for s in slides[2].shapes)               # a picture
    assert sum(1 for s in slides[3].shapes if s.shape_type == 13) == 1     # figure in a column
    assert any(s.has_table for s in slides[4].shapes)                      # a real table
    assert any(s.has_text_frame and "Source: own data" in s.text_frame.text
               for s in slides[1].shapes)


def test_pptx_export_reuses_a_reference_template_layouts(tmp_path):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from NoorSuite import pptxexport as X
    ref = tmp_path / "ref.pptx"
    base = Presentation()                      # stands in for a corporate template
    for lay in base.slide_layouts:             # rename: layouts are found by name
        if lay.name == "Title Slide":
            lay.name = "My title layout"
    base.slides.add_slide(base.slide_layouts[0])           # a sample slide that must disappear
    base.save(str(ref))
    sheets, sheet_png = _sheet_stub(tmp_path)
    src = T.build_document("beamer", "T", ["A"], "", r"\begin{frame}{X}\begin{itemize}\item a"
                           r"\end{itemize}\end{frame}")
    out = tmp_path / "o.pptx"
    ok, _ = X.export_pptx(src, out, sheets, [], sheet_png, lambda *a: None, tmp_path / "w",
                          reference=str(ref), mapping={"title_layout": "My title layout"})
    assert ok
    prs = Presentation(str(out))
    assert [s.slide_layout.name for s in prs.slides] == ["My title layout", "Title and Content"]


def test_pptx_export_through_the_headless_command(tmp_path):
    pytest.importorskip("pptx")
    import subprocess
    import sys
    import pandas as pd
    from NoorSuite.client import OfflineClient
    s = OfflineClient(tmp_path / "p.sciproj")
    s.push_dataframe(pd.DataFrame({"x": [1.0, 2.0], "y": [1.0, 4.0]}), name="d")
    s.plot("d", "x", "y", new_sheet=True, title="Fig A")
    s.push_document("Talk", r"\begin{frame}{Plot}\centering\noorsheet{Fig-A}\end{frame}",
                    template="beamer", title="Talk", authors=["A"])
    out = tmp_path / "talk.pptx"
    r = subprocess.run([sys.executable, "-m", "NoorSuite", "latex", "pptx", "--project",
                        str(tmp_path / "p.sciproj"), "--out", str(out)],
                       capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    from pptx import Presentation
    prs = Presentation(str(out))
    assert len(prs.slides) == 2 and any(sh.shape_type == 13 for sh in prs.slides[1].shapes)


def test_slide_show_navigates_and_covers(app, tmp_path):
    from PyQt6.QtCore import QMarginsF, QSizeF
    from PyQt6.QtGui import QPageLayout, QPageSize, QPainter, QPdfWriter
    from NoorSuite.docview import PresentationView
    pdf = tmp_path / "s.pdf"
    w = QPdfWriter(str(pdf))
    w.setPageSize(QPageSize(QSizeF(160, 100), QPageSize.Unit.Millimeter))
    w.setPageMargins(QMarginsF(0, 0, 0, 0), QPageLayout.Unit.Millimeter)
    p = QPainter(w)
    for i in range(4):
        if i:
            w.newPage()
        p.drawText(100, 100, f"Slide {i + 1}")
    p.end()
    v = PresentationView(str(pdf), 0)
    v.resize(640, 400)
    assert v.page_count() == 4 and v.page == 0
    v.go(1)
    v.go(99)
    assert v.page == 3
    v.go(-5)
    assert v.page == 0
    v._cover = "black"
    v.go(2)
    assert v._cover is None and v.page == 2
    closed = []
    v.closed.connect(closed.append)
    v.close()
    assert closed == [2]
