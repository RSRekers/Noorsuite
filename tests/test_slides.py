"""Free slide objects (\\noorobj), slide setup, the slide canvas and the editor panel."""
import pytest

from NoorSuite import latex
from NoorSuite import slideobjects as so
from tests.test_gui_features import app, win, _doc_win  # noqa: F401  (window fixtures)

TALK = "\n".join([
    r"\documentclass[aspectratio=1610,11pt,t]{beamer}", r"\usepackage{noorsuite}",
    r"\title{T}\author{A}", r"\begin{document}",
    r"\begin{frame}[plain]", r"\titlepage", r"\end{frame}",
    r"\begin{frame}{First}", "Some text", r"\end{frame}",
    r"\begin{frame}{Second}", r"\noorobj[x=1,y=2,w=3,h=1.5,fill=ff8800]{rect}{Hi}",
    r"\noorobj[x=5,y=2,x2=9,y2=4]{arrow}{}", r"\end{frame}",
    r"\end{document}", ""])


def test_objects_roundtrip_and_edits():
    objs = so.objects_in(TALK, 2)
    assert [o["kind"] for o in objs] == ["rect", "arrow"]
    assert objs[0]["fill"] == "ff8800" and objs[0]["w"] == 3.0 and objs[0]["arg"] == "Hi"
    assert objs[1]["x2"] == 9.0
    # only non-default keys are written
    line = so.format_object({"kind": "rect", "x": 1, "y": 2, "w": 3, "h": 1.5, "fill": "ff8800",
                             "arg": "Hi", "line": "000000", "lw": 1.0, "angle": 0})
    assert line == r"\noorobj[x=1,y=2,w=3,h=1.5,fill=ff8800]{rect}{Hi}"
    new = so.replace_object(TALK, 2, 0, {**objs[0], "x": 2.5, "angle": 30})
    assert "x=2.5,y=2,w=3,h=1.5,angle=30,fill=ff8800" in new
    added = so.add_object(TALK, 1, {"kind": "ellipse", "x": 1, "y": 1, "w": 2, "h": 2, "arg": ""})
    s1, e1 = so.frames(added)[1]
    assert len(so.objects_in(added, 1)) == 1 and s1 < added.index(r"\noorobj") < e1
    gone = so.delete_object(TALK, 2, 0)
    assert [o["kind"] for o in so.objects_in(gone, 2)] == ["arrow"]
    moved = so.move_object(TALK, 2, 0, 1)
    assert [o["kind"] for o in so.objects_in(moved, 2)] == ["arrow", "rect"]
    # braces inside the label do not end the object early
    t = so.add_object(TALK, 1, {"kind": "text", "x": 1, "y": 1, "w": 4, "h": 1,
                                "arg": r"\textbf{a} $x^{2}$"})
    assert so.objects_in(t, 1)[0]["arg"] == r"\textbf{a} $x^{2}$"
    assert so.frame_at(TALK, TALK.index("Some text")) == 1
    assert so.slide_size_mm(TALK) == (160.0, 100.0)


def test_slide_setup_line_is_written_replaced_and_removed():
    s = so.write_setup(TALK, {"title-dx": "1", "number-on-title": "true", "body-top": ""})
    assert r"\noorslide{title-dx=1,number-on-title=true}" in s
    assert s.index(r"\noorslide") < s.index(r"\begin{document}")
    assert so.read_setup(s) == {"title-dx": "1", "number-on-title": "true"}
    s2 = so.write_setup(s, {"title-size": "28"})
    assert s2.count(r"\noorslide") == 1 and "title-size=28" in s2 and "title-dx" not in s2
    assert r"\noorslide" not in so.write_setup(s2, {})


def test_canvas_gestures_move_scale_rotate(app):
    from PyQt6.QtCore import QPoint, Qt
    from PyQt6.QtTest import QTest
    from NoorSuite.slidecanvas import SlideCanvas, to_mm
    c = SlideCanvas()
    c.resize(700, 460)
    c.show()
    c.set_state([{"kind": "rect", "x": 4.0, "y": 3.0, "w": 4.0, "h": 2.0, "angle": 0.0,
                  "fill": "ff0000", "arg": ""}], (160.0, 100.0))
    c.snap = False
    got = []
    c.committed.connect(lambda i, o: got.append(o))
    c._layout()
    g = to_mm(c.objs[0])

    def px(x, y):
        p = c.to_px(x, y)
        return QPoint(int(p.x()), int(p.y()))
    # move by 10 mm right / 5 mm down
    QTest.mousePress(c, Qt.MouseButton.LeftButton, pos=px(g["cx"], g["cy"]))
    QTest.mouseMove(c, px(g["cx"] + 10, g["cy"] + 5))
    QTest.mouseRelease(c, Qt.MouseButton.LeftButton, pos=px(g["cx"] + 10, g["cy"] + 5))
    assert got and abs(got[-1]["x"] - 5.0) < 0.05 and abs(got[-1]["y"] - 3.5) < 0.05
    # scale: drag the bottom-right corner 10 mm right
    g = to_mm(c.objs[0])
    corner = (g["cx"] + g["w"] / 2, g["cy"] + g["h"] / 2)
    QTest.mousePress(c, Qt.MouseButton.LeftButton, pos=px(*corner))
    QTest.mouseMove(c, px(corner[0] + 10, corner[1]))
    QTest.mouseRelease(c, Qt.MouseButton.LeftButton, pos=px(corner[0] + 10, corner[1]))
    assert abs(got[-1]["w"] - 5.0) < 0.05 and abs(got[-1]["h"] - 2.0) < 0.05
    assert abs(got[-1]["x"] - 5.0) < 0.05                       # left edge stayed
    # rotate: drag the round handle to the right of the centre -> -90 degrees
    g = to_mm(c.objs[0])
    off = 26.0 / c._scale
    handle = (g["cx"], g["cy"] - g["h"] / 2 - off)
    QTest.mousePress(c, Qt.MouseButton.LeftButton, pos=px(*handle))
    QTest.mouseMove(c, px(g["cx"] + 30, g["cy"]))
    QTest.mouseRelease(c, Qt.MouseButton.LeftButton, pos=px(g["cx"] + 30, g["cy"]))
    assert abs(got[-1]["angle"] + 90.0) < 0.5
    c.close()


def test_editor_panel_edits_the_text(win):
    ids, dm, tab = _doc_win(win, 1)
    tab.editor.setPlainText(TALK)
    ed = tab.slide_editor
    ed.set_frame(1)
    ed.add_object("rect")
    src = tab.source()
    assert len(so.objects_in(src, 1)) == 1 and len(ed.canvas.objs) == 1
    assert ed.canvas.sel == 0
    ed.add_graph() if False else None
    # numeric edit through the property panel
    ed.sp["x"].setValue(3.0)
    assert so.objects_in(tab.source(), 1)[0]["x"] == 3.0
    ed.duplicate_object(0)
    assert len(so.objects_in(tab.source(), 1)) == 2
    ed.reorder_object(1, -1)
    ed.delete_object(0)
    assert len(so.objects_in(tab.source(), 1)) == 1
    # frame navigation + setup dialog values
    ed.set_frame(2)
    assert [o["kind"] for o in ed.canvas.objs] == ["rect", "arrow"]
    tab.apply_source(so.write_setup(tab.source(), {"number-on-title": "true"}))
    assert so.read_setup(tab.source())["number-on-title"] == "true"
    assert tab.editor.toPlainText().count(r"\documentclass") == 1      # minimal edit, no duplication


@pytest.mark.skipif(latex.find_engine() is None, reason="no LaTeX engine installed")
def test_objects_setup_and_figure_offsets_compile(tmp_path):
    src = TALK.replace(r"\title{T}", r"\noorslide{title-dx=0.5,title-size=24,margin-left=1.5,"
                                     r"number-on-title=true,number-x=14,number-y=9}" "\n" r"\title{T}")
    src = src.replace("Some text", r"Some text \noorsheet[width=2cm,xshift=1cm,yshift=2mm]{nothing}")
    src = src.replace(r"\noorobj[x=5,y=2,x2=9,y2=4]{arrow}{}",
                      r"\noorobj[x=5,y=2,x2=9,y2=4,lw=2,dash=dashed]{arrow}{}" "\n"
                      r"\noorobj[x=9,y=5,w=5,h=2,fill=88ccff,fill2=ffffff,grad=radial,line=none,"
                      r"angle=20,opacity=0.7]{ellipse}{$\alpha$ label}" "\n"
                      r"\noorobj[x=1,y=6,w=6,h=1,size=14,align=left]{text}{Free \textbf{text}}")
    latex.prepare_workdir(tmp_path, src)
    res = latex.build_pdf(tmp_path)
    assert res.ok, res.log[-600:]


def test_pptx_export_carries_objects_and_title_number(tmp_path):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from NoorSuite import pptxexport as X
    from NoorSuite.model import SheetModel
    from tests.test_presentation import _png_bytes
    png = _png_bytes(60, 40)

    def sheet_png(_sm, path):
        open(path, "wb").write(png)
    src = TALK.replace(r"\title{T}", r"\noorslide{number-on-title=true,number-x=14,number-y=9}"
                                     "\n" r"\title{T}")
    src = src.replace(r"\noorobj[x=5,y=2,x2=9,y2=4]{arrow}{}",
                      r"\noorobj[x=5,y=2,x2=9,y2=4,lw=2]{arrow}{}" "\n"
                      r"\noorobj[x=2,y=5,w=4,h=3,angle=20,fill=88ccff,fill2=ffffff]{ellipse}{A}" "\n"
                      r"\noorobj[x=8,y=5,w=3,h=2]{sheet}{fig-a}" "\n"
                      r"\noorobj[x=1,y=8,w=6,h=1,size=14]{text}{Free text}")
    P = X.parse_presentation(src)
    assert [f["kind"] for f in P["frames"]] == ["title", "content", "content"]
    assert len(P["frames"][2]["objects"]) == 5 and P["frames"][2]["blocks"] == []
    out = tmp_path / "o.pptx"
    ok, _ = X.export_pptx(src, out, [SheetModel("Fig A", sheet_id="fig-a")], [], sheet_png,
                          lambda *a: None, tmp_path / "w")
    assert ok
    slides = list(Presentation(str(out)).slides)
    kinds = [sh.shape_type for sh in slides[2].shapes]
    assert 13 in kinds                                         # the graph picture
    assert any("Connector" in sh.name for sh in slides[2].shapes)         # the arrow
    assert any(sh.has_text_frame and sh.text_frame.text == "Free text" for sh in slides[2].shapes)
    assert any(abs(sh.rotation - (-20)) < 0.5 or abs(sh.rotation - 340) < 0.5
               for sh in slides[2].shapes if hasattr(sh, "rotation"))
    assert "slidenum" in slides[0]._element.xml                 # number on the title slide


def test_figure_call_options_are_edited_in_place():
    from NoorSuite import figureopts as F
    src = ("Text \\noorsheet[width=0.8\\textwidth,keepaspectratio]{fig-a} and\n"
           "% \\noorsheet{commented}\n\\noorfile{logo}\n")
    calls = F.find_calls(src)
    assert [(c["cmd"], c["key"], c["line"]) for c in calls] == [
        ("noorsheet", "fig-a", 0), ("noorfile", "logo", 2)]
    new = F.replace_call(src, calls[0], {"width": "6cm", "angle": "15", "xshift": "1cm",
                                         "height": ""})
    assert r"\noorsheet[keepaspectratio,width=6cm,angle=15,xshift=1cm]{fig-a}" in new
    plain = F.replace_call(src, calls[1], {"scale": "0.5"})
    assert r"\noorfile[scale=0.5]{logo}" in plain
    assert F.get_options("a=1, b={x,y}, flag") == {"a": "1", "b": "{x,y}", "flag": ""}


def test_transform_dialog_reads_and_writes_values(app):
    from NoorSuite import figureopts as F
    call = F.find_calls(r"\noorsheet[width=0.5\textwidth,angle=10]{k}")[0]
    d = F.FigureTransformDialog(call)
    assert d.unit.currentIndex() == 1 and abs(d.width.value() - 50) < 1e-6
    assert d.angle.value() == 10
    d.unit.setCurrentIndex(0)
    d.width.setValue(7.5)
    d.xshift.setValue(-0.5)
    up = d.updates()
    assert up["width"] == "7.5cm" and up["xshift"] == "-0.5cm" and up["angle"] == "10"
    assert up["height"] == "" and up["scale"] == ""


def test_text_formatting_wraps_the_selection(win):
    from NoorSuite import textformat as TF
    assert TF.wrap("bold", "x") == r"\textbf{x}"
    assert TF.wrap("color", "x", "#c0392b") == r"\textcolor[HTML]{C0392B}{x}"
    assert TF.wrap("size", "x", "Large") == r"{\Large x}"
    assert TF.wrap("itemize", "a\nb") == "\\begin{itemize}\n  \\item a\n  \\item b\n\\end{itemize}"
    assert TF.wrap("sup", "2") == "$^{2}$"
    ids, dm, tab = _doc_win(win, 1)
    tab.editor.setPlainText("hello world")
    cur = tab.editor.textCursor()
    cur.setPosition(0)
    cur.setPosition(5, cur.MoveMode.KeepAnchor)
    tab.editor.setTextCursor(cur)
    tab.format_selection("bold")
    assert tab.editor.toPlainText() == r"\textbf{hello} world"
    tab.editor.undo()
    assert tab.editor.toPlainText() == "hello world"


CLEAN = "\n".join([
    r"\documentclass[aspectratio=1610,11pt,t]{beamer}", r"\usetheme{noorclean}",
    r"\usepackage{noorsuite}", r"\title{T}\author{A}\date{D}", r"\begin{document}",
    r"\begin{frame}[plain]\titlepage\end{frame}",
    r"\begin{frame}{Results here}", r"\begin{itemize}\item First point\end{itemize}",
    r"\framesource{Source: me}", r"\end{frame}", r"\end{document}", ""])


def test_master_objects_and_placeholders_in_the_text():
    src = so.add_object(CLEAN, so.MASTER, {"kind": "rect", "x": 0, "y": 9.9, "w": 16, "h": 0.1,
                                           "fill": "C8553D", "line": "none", "arg": ""})
    assert r"\noormasterobj[x=0,y=9.9,w=16,h=0.1,fill=C8553D,line=none]{rect}{}" in src
    assert src.index(r"\noormasterobj") < src.index(r"\begin{document}")
    objs = so.objects_in(src, so.MASTER)
    assert len(objs) == 1 and objs[0]["master"] and objs[0]["fill"] == "C8553D"
    assert so.objects_in(src, 1) == []                               # frames are separate
    src = so.delete_object(src, so.MASTER, 0)
    assert r"\noormasterobj" not in src
    phs = so.placeholders(CLEAN, (16.0, 10.0))
    assert [p["kind"] for p in phs] == ["ph-title", "ph-body", "ph-number", "ph-source"]
    phs[0]["x"], phs[0]["y"], phs[0]["w"], phs[0]["size"] = 2.0, 1.0, 11.0, 24.0
    phs[2]["x"], phs[2]["y"], phs[2]["color"] = 14.0, 0.4, "FFFFFF"
    out = so.write_placeholders(CLEAN, phs)
    setup = so.read_setup(out)
    assert setup["title-x"] == "2" and setup["title-size"] == "24" and setup["num-color"] == "FFFFFF"
    assert so.has_placeholders(out)
    again = so.placeholders(out, (16.0, 10.0))
    assert again[0]["x"] == 2.0 and again[2]["color"] == "FFFFFF"


@pytest.mark.skipif(latex.find_engine() is None, reason="no LaTeX engine installed")
def test_master_layout_keys_compile_and_are_detected(win, tmp_path):
    ids, dm, tab = _doc_win(win, 1)
    src = so.write_placeholders(CLEAN, [
        {**p, "x": 2.0 if p["kind"] == "ph-title" else p["x"]} for p in so.placeholders(
            CLEAN, (16.0, 10.0))])
    src = so.add_object(src, so.MASTER, {"kind": "ellipse", "x": 14, "y": 0.2, "w": 1.5, "h": 0.6,
                                         "fill": "88ccff", "arg": ""})
    latex.prepare_workdir(tmp_path, src)
    res = latex.build_pdf(tmp_path)
    assert res.ok, res.log[-800:]
    # the editor measures the theme's own layout in the compiled PDF
    tab.editor.setPlainText(CLEAN)
    latex.prepare_workdir(tab.workdir, CLEAN)
    res = latex.build_pdf(tab.workdir)
    assert res.ok, res.log[-800:]
    tab._pdf = res.pdf
    tab.built_source = CLEAN
    ed = tab.slide_editor
    ed.set_frame(so.MASTER)
    kinds = [o["kind"] for o in ed.canvas.objs]
    assert kinds == ["ph-title", "ph-body", "ph-number", "ph-source"]
    title = ed.canvas.objs[0]
    assert 0.5 < title["x"] < 1.5 and title["y"] < 2.0           # noorclean: x 0.9, bar at the top
    # moving a placeholder through the panel writes the layout line and keeps the others
    ed.canvas.select(0)
    ed.sp["x"].setValue(3.0)
    ed.sp["x"].editingFinished.emit() if False else None
    setup = so.read_setup(tab.source())
    assert setup.get("title-x") == "3" and "num-x" in setup and "src-x" in setup
    ed.add_object("rect")
    assert any(o.get("master") for o in so.objects_in(tab.source(), so.MASTER))


@pytest.mark.skipif(latex.find_engine() is None or "tectonic" not in str(latex.find_engine()).lower(),
                    reason="needs tectonic")
def test_second_build_is_a_single_tex_pass(tmp_path):
    first = CLEAN.replace("First point", r"First \noorobj[x=1,y=5,w=2,h=1]{rect}{x}")
    latex.prepare_workdir(tmp_path, first)
    r1 = latex.build_pdf(tmp_path)
    assert r1.ok and r1.log.count("Running TeX") == 2                 # positions need a second run
    latex.prepare_workdir(tmp_path, first.replace("Results here", "Other title"))
    r2 = latex.build_pdf(tmp_path)
    assert r2.ok and r2.log.count("Running TeX") == 1                 # aux of the last build reused
    assert (tmp_path / "out" / "main.synctex.gz").is_file()
