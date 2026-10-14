"""Overlay effects (pause / item<2-> / uncover), PowerPoint animations, find & replace."""
import re

import pytest

from NoorSuite import findbar
from NoorSuite import pptxexport as X
from tests.test_gui_features import app, win  # noqa: F401  (window fixtures)
from tests.test_presentation import _png_bytes


def _frame(body):
    return (r"\documentclass{beamer}" "\n" r"\begin{document}" "\n" r"\begin{frame}{T}" "\n"
            + body + "\n" r"\end{frame}" "\n" r"\end{document}")


def test_overlay_steps_of_pause_items_and_uncover():
    body = (r"\begin{itemize}[<+->]" "\n" r"\item One" "\n" r"\item Two" "\n" r"\end{itemize}" "\n"
            "Always here\n" r"\pause" "\n" "After pause\n"
            r"\uncover<5->{\noorsheet{a}}" "\n"
            r"\begin{itemize}\item A\item<7-> B\end{itemize}")
    f = X.parse_presentation(_frame(body))["frames"][0]
    assert f["blocks"][0][2] == [1, 2]                          # <+-> : one click per item
    assert f["blocks"][1] == ("text", "Always here")            # not hidden by the list
    assert f["blocks"][2] == ("text", "After pause", 3)         # pause continues the counter
    assert f["blocks"][3][-1] == 5 and f["blocks"][3][0] == "picture"
    assert f["blocks"][4][2] == [3, 7]


def test_transition_commands_are_read():
    f = X.parse_presentation(_frame(r"\transdissolve" "\n" r"\begin{itemize}\item a\end{itemize}")
                             )["frames"][0]
    assert f["transition"] == "dissolve"
    src = (r"\documentclass{beamer}\transwipe\begin{document}\begin{frame}{A}x\end{frame}"
           r"\begin{frame}{B}\transfade y\end{frame}\end{document}")
    fr = X.parse_presentation(src)["frames"]
    assert [x["transition"] for x in fr] == ["wipe", "fade"]


def test_pptx_gets_click_animations_and_transition(tmp_path):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from NoorSuite.model import SheetModel
    png = _png_bytes(60, 40)
    sm = SheetModel("Fig A", sheet_id="fig-a")

    def sheet_png(_sm, path):
        open(path, "wb").write(png)
    body = (r"\transfade" "\n" r"\begin{itemize}[<+->]\item One\item Two\end{itemize}" "\n"
            r"\pause" "\n" r"\noorsheet{fig-a}")
    out = tmp_path / "e.pptx"
    ok, _ = X.export_pptx(_frame(body), out, [sm], [], sheet_png, lambda *a: None,
                          tmp_path / "w", effect="fade")
    assert ok
    slide = list(Presentation(str(out)).slides)[0]
    xml = slide._element.xml
    assert xml.count('nodeType="clickEffect"') == 3             # item 1, item 2, the picture
    assert xml.count('<p:pRg') == 4                              # set + animEffect per paragraph
    assert 'filter="fade"' in xml and "<p:transition" in xml
    assert xml.index("<p:transition") < xml.index("<p:timing")
    # same-step content appears together: appear effect has no animEffect
    out2 = tmp_path / "e2.pptx"
    X.export_pptx(_frame(body), out2, [sm], [], sheet_png, lambda *a: None, tmp_path / "w2")
    assert "animEffect" not in list(Presentation(str(out2)).slides)[0]._element.xml


def test_find_matches_options_and_replace_all():
    t = "Cat cat concatenate CAT"
    assert findbar.find_matches(t, "cat") == [(0, 3), (4, 7), (11, 14), (20, 23)]
    assert findbar.find_matches(t, "cat", case=True) == [(4, 7), (11, 14)]
    assert findbar.find_matches(t, "cat", word=True) == [(0, 3), (4, 7), (20, 23)]
    assert findbar.find_matches(t, r"c\w+t", regex=True)[0] == (0, 3)
    new, n = findbar.replace_all_text(r"\ref{a} \ref{b}", r"\\ref\{(\w)\}", r"\\autoref{\1}",
                                      regex=True)
    assert (new, n) == (r"\autoref{a} \autoref{b}", 2)
    assert findbar.replace_all_text("a.b a.b", ".", "-")[0] == "a-b a-b"      # literal dot


def test_find_bar_in_the_document_tab(win):
    from PyQt6.QtGui import QTextCursor
    win.sheets  # the window exists
    from tests.test_gui_features import _doc_win
    _ids, _dm, tab = _doc_win(win, 1)
    tab.editor.setPlainText("alpha beta alpha\nGamma alpha")
    fb = tab.findbar
    fb.open_find(False)
    fb.find_edit.setText("alpha")
    assert fb.count_label.text() in ("1 of 3", "3 matches")
    fb.step(False)
    a = tab.editor.textCursor()
    assert a.selectedText() == "alpha"
    start = a.selectionStart()
    fb.step(False)
    assert tab.editor.textCursor().selectionStart() > start
    fb.replace_edit.setText("ALPHA")
    n = fb.replace_all()
    assert n == 3 and tab.editor.toPlainText() == "ALPHA beta ALPHA\nGamma ALPHA"
    tab.editor.undo()                                           # one undo step restores all
    assert tab.editor.toPlainText() == "alpha beta alpha\nGamma alpha"
    fb.find_edit.setText("(")                                    # plain text: no regex error
    assert fb.count_label.text() in ("no match",)
    fb.regex_check.setChecked(True)
    assert fb.count_label.text() == "invalid pattern"
    fb.close_bar()
    assert not fb.isVisible()
    assert isinstance(QTextCursor, type)


def test_effects_menu_edits_the_source(win):
    from tests.test_gui_features import _doc_win
    _ids, _dm, tab = _doc_win(win, 1)
    tab.editor.setPlainText("\\documentclass{beamer}\n\\begin{document}\n"
                            "\\begin{frame}{x}\n\\begin{itemize}\n\\item a\n\\end{itemize}\n"
                            "\\end{frame}\n\\end{document}")
    cur = tab.editor.textCursor()
    cur.setPosition(tab.editor.toPlainText().index("\\item a"))
    tab.editor.setTextCursor(cur)
    tab.reveal_list()
    assert "\\begin{itemize}[<+->]" in tab.editor.toPlainText()
    tab.add_preamble_line("\\setbeamercovered{transparent}")
    tab.add_preamble_line("\\setbeamercovered{transparent}")      # only once
    src = tab.editor.toPlainText()
    assert src.count("setbeamercovered") == 1
    assert src.index("setbeamercovered") < src.index("\\begin{document}")


def test_slide_show_fade_flag(app, tmp_path):
    from PyQt6.QtCore import QMarginsF, QSizeF
    from PyQt6.QtGui import QPageLayout, QPageSize, QPainter, QPdfWriter
    from NoorSuite.docview import PresentationView
    pdf = tmp_path / "s.pdf"
    w = QPdfWriter(str(pdf))
    w.setPageSize(QPageSize(QSizeF(160, 100), QPageSize.Unit.Millimeter))
    w.setPageMargins(QMarginsF(0, 0, 0, 0), QPageLayout.Unit.Millimeter)
    p = QPainter(w)
    p.drawText(100, 100, "one")
    w.newPage()
    p.drawText(100, 100, "two")
    p.end()
    v = PresentationView(str(pdf), 0, fade=True)
    v.resize(320, 200)
    v.show()
    v.go(1)
    assert v.page == 1 and v._blend <= 1.0
    v.close()
